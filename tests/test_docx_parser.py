from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from backend.parser import DocumentError, parse_document


def test_docx_parser_exact_text_inline_numbers_and_letters(make_docx):
    raw = make_docx(["3.4. Подразделения:", "а. Отдел анализа.",
                     "3.9. Первая норма. 3.10.Вторая норма. 3.11. Третья норма."])
    doc = parse_document(raw, "example.docx")
    assert [c.section for c in doc.clauses] == ["3.4", "3.4.а", "3.9", "3.10", "3.11"]
    assert doc.clauses[1].text == "а. Отдел анализа."
    assert "".join(c.text for c in doc.clauses[2:]) == "3.9. Первая норма. 3.10.Вторая норма. 3.11. Третья норма."
    assert parse_document(raw, "example.docx").clauses == doc.clauses


def test_tables_in_document_order():
    doc = Document()
    doc.add_paragraph("1. Первая норма.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2. Норма в таблице."
    table.cell(0, 1).text = "3. Другая ячейка."
    doc.add_paragraph("4. Последняя норма.")
    stream = BytesIO()
    doc.save(stream)
    parsed = parse_document(stream.getvalue())
    assert [c.section for c in parsed.clauses] == ["1", "2", "3", "4"]


def test_word_numbering_and_restart():
    doc = Document()
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "99")
    for level in range(2):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1." if level == 0 else "%1.%2.")):
            item = OxmlElement("w:" + tag)
            item.set(qn("w:val"), value)
            lvl.append(item)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "99")
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), "99")
    num.append(ref)
    numbering.append(num)
    for level in (0, 1, 1, 0, 1):
        paragraph = doc.add_paragraph("Нумерация Word без номера в тексте.")
        props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        props.get_or_add_numId().val = 99
        props.get_or_add_ilvl().val = level
    stream = BytesIO()
    doc.save(stream)
    parsed = parse_document(stream.getvalue())
    assert [c.section for c in parsed.clauses] == ["1", "1.1", "1.2", "2", "2.1"]
    assert all(c.text == "Нумерация Word без номера в тексте." for c in parsed.clauses)


def test_unnumbered_text_does_not_invent_section(make_docx):
    parsed = parse_document(make_docx(["Общие положения без номера."]))
    assert parsed.clauses[0].section is None


@pytest.mark.parametrize("raw,name", [(b"not a zip", "bad.docx"), (b"", "file.pdf")])
def test_invalid_upload(raw, name):
    with pytest.raises(DocumentError):
        parse_document(raw, name)


def test_flattened_contents_is_not_extracted_twice(make_docx):
    doc = parse_document(make_docx(["1. Основные положения", "1.1. Норма.", "Содержание", "1. Основные положения 3"]))
    assert [c.section for c in doc.clauses] == ["1", "1.1"]


@pytest.mark.parametrize("revision", [8, 9])
def test_real_local_documents(revision):
    paths = list((Path(__file__).resolve().parents[1] / "data").glob(f"*_{revision}_*.docx"))
    if not paths:
        pytest.skip("Local private DOCX is deliberately not committed.")
    parsed = parse_document(paths[0])
    assert len(parsed.clauses) > 100
    assert any(c.section == "3.4" for c in parsed.clauses)
    assert any(c.section == "5.4.1" for c in parsed.clauses)
    assert all(c.text.strip() for c in parsed.clauses)
